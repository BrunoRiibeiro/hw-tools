#!/usr/bin/env python3

import argparse
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_border(cell, border_color="000000", border_size="4"):
	tc = cell._element
	tcPr = tc.get_or_add_tcPr()
	for border in tcPr.xpath('.//w:TCBorders'):
		tcPr.remove(border)
	border_xml = f'''
	<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
		<w:top w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
		<w:left w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
		<w:bottom w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
		<w:right w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
	</w:tcBorders>
	'''
	tcPr.append(parse_xml(border_xml))

parser = argparse.ArgumentParser(description="Process command-line arguments for the script.")
parser.add_argument('wetask_table', help='Input XLSX file from WeTask')
parser.add_argument('sr_table', help='Input XLSX file with SRs')
parser.add_argument('--out-file', default="output.docx", help='Output file name (default: output.docx)')
args = parser.parse_args()

wetask_table = args.wetask_table
sr_table = args.sr_table
output_docx = args.out_file

df = pd.read_excel(wetask_table)
dz = pd.read_excel(sr_table)
SRs = 0
SR_SERVICE = 1
SR_NAME = 2
SR_NUMB = 5
SR_DESCRIPTION = 10
SR_CATEGORY = 12
SR_SOLUTION = 15

doc = Document()
table = doc.add_table(rows=1, cols=6)
tab_headers = [ "Ticket", "Issue", "Ticket Information", "Solution", "Service", "Category"]
for i, header in enumerate(tab_headers):
	table.cell(0, i).paragraphs[0].add_run(header).bold = True
	table.cell(0, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

def to_text(val):
	if pd.isna(val):
		return ""
	return str(val)

wetask_map = {}
for r in df.itertuples(index=False):
	sr_key = r[SR_NUMB]
	if pd.isna(sr_key):
		continue

	wetask_map[sr_key] = {
			"Issue": r[SR_NAME],
			"Ticket Information": r[SR_DESCRIPTION],
			"Solution": r[SR_SOLUTION],
			"Service": r[SR_SERVICE],
			"Category": r[SR_CATEGORY],
			}

# Monta as linhas no DOCX a partir da tabela closed
for r in dz.itertuples(index=False):
	ticket_sr = r[SRs]
	if pd.isna(ticket_sr):
		continue

	wetask_data = wetask_map.get(ticket_sr)
	new_cells = table.add_row().cells
	new_cells[0].text = to_text(ticket_sr)
	new_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

	if wetask_data:
		new_cells[1].text = to_text(wetask_data["Issue"])
		new_cells[2].text = to_text(wetask_data["Ticket Information"])
		new_cells[3].text = to_text(wetask_data["Solution"])
		new_cells[4].text = to_text(wetask_data["Service"])
		new_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
		new_cells[5].text = to_text(wetask_data["Category"])
		new_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

for row in table.rows:
	for cell in row.cells:
		set_cell_border(cell)

for row in table.rows:
	for cell in row.cells:
		for paragraph in cell.paragraphs:
			for run in paragraph.runs:
				run.font.name = 'Times New Roman'
				run.font.size = Pt(11)

doc.save(output_docx)
print(f"File {output_docx} created!")
