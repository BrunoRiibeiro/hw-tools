#!/usr/bin/env python3

import argparse
import sys
import pandas as pd
from datetime import datetime
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH

#             "Raul"      "Cassio"    "Gabriel"   "Andre"
HUAWEI_IDs = ["50054635", "00959645", "00934972", "50057467"]

def set_cell_border(cell, border_color="000000", border_size="4"):
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    for border in tcPr.xpath('.//w:TCBorders'):
        tcPr.remove(border)
    border_xml = f'''
    <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:top w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
        <w:left w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
        <w:bottom w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
        <w:right w:val="single" w:sz="{border_size}" w:color="{border_color}"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(border_xml))

parser = argparse.ArgumentParser(description="Process command-line arguments for the script.")
parser.add_argument('srs_file', help='Input XLSX file')
parser.add_argument('sline_srs', type=int, help='Starting line index')
parser.add_argument('eline_srs', type=int, help='Ending line index')
parser.add_argument('daily_file', help='Input XLSX file')
parser.add_argument('sline_daily', type=int, help='Starting line index')
parser.add_argument('eline_daily', type=int, help='Ending line index')
parser.add_argument('--acc', type=int, default=1, help='Value of acc (default: 1)')
parser.add_argument('--out-file', default="output.docx", help='Output file name (default: output.docx)')
args = parser.parse_args()

srs_file = args.srs_file
sline_srs = args.sline_srs
eline_srs = args.eline_srs
daily_file = args.daily_file
sline_daily = args.sline_daily
eline_daily = args.eline_daily
acc = args.acc
output_docx = args.out_file

df = pd.read_excel(srs_file)
dz = pd.read_excel(daily_file)
SR = 1
SR_NAME = 3
SR_STATUS = 4
SR_START_DATE = 5
SR_DESCRIPTION = 7
SR_OWNER = 12
DL_NAME = 1
DL_STATUS = 2
DL_START_DATE = 3
DL_DESCRIPTION = 5
DL_OWNER = 12

doc = Document()
table = doc.add_table(rows=1, cols=3)
tab_headers = [ "No.", "This week's work", "Remarks"]
for i, header in enumerate(tab_headers):
	table.cell(0, i).paragraphs[0].add_run(header).bold = True
	table.cell(0, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

tasks_list = []
for index, row in df.iloc[sline_srs-1:eline_srs].iterrows():
	#	print(row.iloc[SR_OWNER])
	#	print(row.iloc[SR_OWNER].split()[-1])
	#	if row.iloc[SR_OWNER].split()[-1] not in HUAWEI_IDs:
	sr = None if pd.isna(row.iloc[SR]) else int(row.iloc[SR])
	tasks_list.append([pd.to_datetime(row.iloc[SR_START_DATE]), sr,
					row.iloc[SR_NAME], row.iloc[SR_STATUS],
					row.iloc[SR_DESCRIPTION]])

for index, row in dz.iloc[sline_daily-1:eline_daily].iterrows():
	tasks_list.append([pd.to_datetime(row.iloc[DL_START_DATE]), "no",
					row.iloc[DL_NAME], row.iloc[DL_STATUS],
					row.iloc[DL_DESCRIPTION]])

stasks_list = sorted(tasks_list, key=lambda x: x[0])
for row in stasks_list:
	start_date = row[0].strftime("%d.%m.%y")
	sr = row[1]
	name = row[2]
	status = row[3]
	description = row[4]
	text_intro = f"On {start_date} - (SR {sr}) - {name}."
	text_desc = f"\n\n{description}\n"

	row_cells = table.add_row().cells
	row_cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
	num_cell = row_cells[0].paragraphs[0]
	num_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
	num_cell.add_run(str(acc))

	desc_cell = row_cells[1].paragraphs[0]
	desc_cell.add_run(text_intro).bold = True
	desc_cell.add_run(text_desc)

	row_cells[2].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
	status_cell = row_cells[2].paragraphs[0]
	status_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
	status_cell.add_run(status).italic = True

	acc += 1

for row in table.rows:
	for cell in row.cells:
		set_cell_border(cell)

for row in table.rows:
	for cell in row.cells:
		for paragraph in cell.paragraphs:
			for run in paragraph.runs:
				run.font.name = 'SimSun'
				run.font.size = Pt(11)

doc.save(output_docx)
print(f"File {output_docx} created!")
